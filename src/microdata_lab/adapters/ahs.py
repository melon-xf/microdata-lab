from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import httpx
import numpy as np
import pandas as pd
import pyarrow.csv as pacsv
import pyarrow.parquet as pq
import yaml
from openpyxl import load_workbook
from pydantic import AnyHttpUrl, BaseModel, ConfigDict, TypeAdapter

from microdata_lab.adapters.base import SourceAdapter
from microdata_lab.models import (
    DiscoveredArtifact,
    DiscoveredRelease,
    StoredArtifact,
    ValidationResult,
)

NATIONAL_CSV = "national_csv"
VERSION_CONTROL = "version_control"
VERIFICATION_WORKBOOK = "verification_workbook"
VALUE_LABELS = "value_labels"
DESIGN_REPORT = "design_report"
VARIANCE_GUIDE = "variance_guide"
CODEBOOK_VARIABLES = "codebook_variables"
CODEBOOK_DETAILS = "codebook_details"
CODEBOOK_YEARS = "codebook_years"
GETTING_STARTED = "getting_started"

_http_url = TypeAdapter(AnyHttpUrl)
_DEFINITION_PATH = Path(__file__).resolve().parents[3] / "config" / "ahs" / "2023.yaml"


class AHSStaticArtifact(BaseModel):
    model_config = ConfigDict(extra="forbid")

    role: str
    url: AnyHttpUrl
    filename: str
    documentation: bool = False


class AHSTableContract(BaseModel):
    model_config = ConfigDict(extra="forbid")

    filename: str
    rows: int
    columns: int
    key: list[str]


class AHSWeightContract(BaseModel):
    model_config = ConfigDict(extra="forbid")

    full_sample: str
    split_sample: list[str]
    replicate_prefix: str
    replicate_count: int
    variance_factor: float


class AHSBenchmark(BaseModel):
    model_config = ConfigDict(extra="forbid")

    table: str
    universe_variable: str
    universe_values: list[str]
    expected_status_values: list[str]
    expected_estimate_thousands: int
    expected_moe90_thousands: int
    estimate_rounding_tolerance: float


class AHSCodebookContract(BaseModel):
    model_config = ConfigDict(extra="forbid")

    minimum_variables: int
    minimum_details: int
    required_year_option: str
    required_variables: list[str]


class AHSReleaseDefinition(BaseModel):
    model_config = ConfigDict(extra="forbid")

    year: int
    revision: str
    landing_page: AnyHttpUrl
    required_artifacts: list[str]
    static_artifacts: list[AHSStaticArtifact]
    tables: dict[str, AHSTableContract]
    weight_contract: AHSWeightContract
    benchmark: AHSBenchmark
    codebook_contract: AHSCodebookContract


class AHSAdapter(SourceAdapter):
    slug = "ahs"

    def __init__(
        self,
        *,
        definition: AHSReleaseDefinition | None = None,
        client: httpx.Client | None = None,
    ) -> None:
        self.definition = definition or _load_definition(_DEFINITION_PATH.parent)
        self._owns_client = client is None
        self.client = client or httpx.Client(
            follow_redirects=True,
            timeout=httpx.Timeout(120.0, connect=30.0),
            headers={"User-Agent": "microdata-lab/0.1 (public microdata research)"},
        )

    def close(self) -> None:
        if self._owns_client:
            self.client.close()

    def available_years(self) -> list[int]:
        response = self.client.get(str(self.definition.landing_page))
        response.raise_for_status()
        marker = f"AHS {self.definition.year} National PUF v{self.definition.revision} CSV"
        return [self.definition.year] if marker in response.text else []

    def discover(self, year: int | None = None) -> DiscoveredRelease:
        selected = self.definition.year if year is None else year
        if selected != self.definition.year:
            raise ValueError(f"AHS release definition not pinned for year {selected}")

        response = self.client.get(str(self.definition.landing_page))
        response.raise_for_status()
        links = _extract_links(response.text, str(self.definition.landing_page))
        page_artifacts = _resolve_page_artifacts(
            links,
            year=selected,
            revision=self.definition.revision,
        )
        artifacts = list(page_artifacts)
        artifacts.extend(
            DiscoveredArtifact(
                role=item.role,
                url=item.url,
                filename=item.filename,
                documentation=item.documentation,
                link_text=item.filename,
            )
            for item in self.definition.static_artifacts
        )
        roles = {_role_value(item.role) for item in artifacts}
        missing = set(self.definition.required_artifacts) - roles
        if missing:
            raise RuntimeError(f"AHS discovery missing required roles: {sorted(missing)}")

        return DiscoveredRelease(
            survey=self.slug,
            year=selected,
            landing_page=self.definition.landing_page,
            artifacts=artifacts,
            source_metadata={
                "revision": self.definition.revision,
                "sample": "Integrated National Sample",
                "record_structure": "relational",
                "record_unit": "housing unit, person, mortgage, or project by table",
                "universe": "residential housing units in the 50 states and DC",
                "full_sample_weight": self.definition.weight_contract.full_sample,
                "split_sample_weights": self.definition.weight_contract.split_sample,
                "replicate_count": self.definition.weight_contract.replicate_count,
                "variance_factor": self.definition.weight_contract.variance_factor,
                "merge_key": "CONTROL",
                "benchmark": self.definition.benchmark.model_dump(),
                "codebook_url": (
                    "https://www.census.gov/data-tools/demo/codebook/ahs/ahsdict.html"
                ),
                "release_status": "official national PUF revision 1.1",
            },
        )

    def validate_release(
        self,
        run_root: Path,
        release: DiscoveredRelease,
        artifacts: list[StoredArtifact],
    ) -> ValidationResult:
        roles = {_role_value(artifact.role) for artifact in artifacts}
        required = set(self.definition.required_artifacts)
        checks = {"ahs_required_artifacts_present": required <= roles}
        extracted = _extracted_tables(run_root, artifacts, self.definition)
        checks["ahs_exact_relational_table_set"] = set(extracted) == set(self.definition.tables)

        headers: dict[str, list[str]] = {}
        household_controls: set[str] = set()
        table_rows: dict[str, int] = {}
        for table_name, contract in self.definition.tables.items():
            path = extracted.get(table_name)
            if path is None:
                checks[f"ahs_{table_name}_present"] = False
                continue
            header = _csv_header(path)
            headers[table_name] = header
            checks[f"ahs_{table_name}_columns"] = len(header) == contract.columns
            rows, _keys, controls, duplicates, null_keys = _validate_relational_table(
                path,
                contract,
            )
            table_rows[table_name] = rows
            checks[f"ahs_{table_name}_rows"] = rows == contract.rows
            checks[f"ahs_{table_name}_keys_nonmissing"] = not null_keys
            if contract.key:
                checks[f"ahs_{table_name}_keys_unique"] = not duplicates
            if table_name == "household":
                household_controls = controls
            else:
                checks[f"ahs_{table_name}_foreign_keys"] = controls <= household_controls

        household = extracted.get("household")
        if household is None:
            return ValidationResult(passed=False, checks=checks, notes=[])

        weight_contract = self.definition.weight_contract
        replicate_names = [
            f"{weight_contract.replicate_prefix}{index}"
            for index in range(1, weight_contract.replicate_count + 1)
        ]
        household_header = headers.get("household", [])
        required_weight_columns = [
            weight_contract.full_sample,
            *weight_contract.split_sample,
            *replicate_names,
        ]
        checks["ahs_weight_columns_complete"] = all(
            name in household_header for name in required_weight_columns
        )
        checks["ahs_replicate_count"] = (
            len(
                [
                    name
                    for name in household_header
                    if re.fullmatch(rf"{re.escape(weight_contract.replicate_prefix)}\d+", name)
                ]
            )
            == weight_contract.replicate_count
        )

        benchmark = self.definition.benchmark
        read_columns = [
            benchmark.universe_variable,
            weight_contract.full_sample,
            *weight_contract.split_sample,
            *replicate_names,
        ]
        full_total = 0.0
        replicate_totals = np.zeros(weight_contract.replicate_count, dtype=float)
        weights_valid = True
        split_weights_valid = True
        statuses: set[str] = set()
        for chunk in pd.read_csv(
            household,
            usecols=read_columns,
            chunksize=5_000,
            low_memory=False,
        ):
            status = chunk[benchmark.universe_variable].map(_normalize_code)
            statuses.update(status.dropna().astype(str).unique())
            full_weights = pd.to_numeric(
                chunk[weight_contract.full_sample], errors="coerce"
            ).to_numpy(dtype=float)
            replicate_values = (
                chunk[replicate_names].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
            )
            split_values = (
                chunk[weight_contract.split_sample]
                .apply(pd.to_numeric, errors="coerce")
                .to_numpy(dtype=float)
            )
            weights_valid &= bool(
                np.isfinite(full_weights).all()
                and (full_weights > 0).all()
                and np.isfinite(replicate_values).all()
                and (replicate_values > 0).all()
            )
            split_weights_valid &= bool(
                np.isfinite(split_values).all() and (split_values >= 0).all()
            )
            mask = status.isin(benchmark.universe_values).to_numpy(dtype=bool)
            full_total += float(full_weights[mask].sum())
            replicate_totals += replicate_values[mask].sum(axis=0)

        estimate_thousands = full_total / 1000.0
        replicate_estimates = replicate_totals / 1000.0
        standard_error = math.sqrt(
            weight_contract.variance_factor
            * float(np.square(replicate_estimates - estimate_thousands).sum())
        )
        moe90 = 1.645 * standard_error
        checks.update(
            {
                "ahs_weights_positive_and_finite": weights_valid,
                "ahs_split_weights_nonnegative_and_finite": split_weights_valid,
                "ahs_interview_status_domain": statuses == set(benchmark.expected_status_values),
                "ahs_benchmark_estimate": math.isclose(
                    estimate_thousands,
                    float(benchmark.expected_estimate_thousands),
                    abs_tol=benchmark.estimate_rounding_tolerance,
                    rel_tol=0.0,
                ),
                "ahs_benchmark_moe90": math.isclose(
                    moe90,
                    float(benchmark.expected_moe90_thousands),
                    abs_tol=benchmark.estimate_rounding_tolerance,
                    rel_tol=0.0,
                ),
            }
        )

        verification = _artifact_path(run_root, artifacts, VERIFICATION_WORKBOOK)
        workbook = load_workbook(verification, data_only=True, read_only=True)
        worksheet = workbook["General Housing"]
        checks["ahs_verification_workbook_contract"] = (
            worksheet["B6"].value == "INTSTATUS = '1'"
            and int(worksheet["D6"].value) == benchmark.expected_estimate_thousands
            and int(worksheet["F6"].value) == benchmark.expected_moe90_thousands
        )
        workbook.close()

        variables = json.loads(_artifact_path(run_root, artifacts, CODEBOOK_VARIABLES).read_text())
        details = json.loads(_artifact_path(run_root, artifacts, CODEBOOK_DETAILS).read_text())
        years = json.loads(_artifact_path(run_root, artifacts, CODEBOOK_YEARS).read_text())
        codebook = self.definition.codebook_contract
        variable_names = {str(item.get("summary", {}).get("name")) for item in variables}
        checks.update(
            {
                "ahs_codebook_variable_count": len(variables) >= codebook.minimum_variables,
                "ahs_codebook_detail_count": len(details) >= codebook.minimum_details,
                "ahs_codebook_required_variables": set(codebook.required_variables)
                <= variable_names,
                "ahs_codebook_year": any(
                    str(item.get("id")) == codebook.required_year_option for item in years
                ),
            }
        )

        notes = [
            f"household_rows={table_rows.get('household', 0)}",
            f"mortgage_rows={table_rows.get('mortgage', 0)}",
            f"person_rows={table_rows.get('person', 0)}",
            f"project_rows={table_rows.get('project', 0)}",
            f"benchmark_estimate_thousands={estimate_thousands:.9f}",
            f"benchmark_se_thousands={standard_error:.9f}",
            f"benchmark_moe90_thousands={moe90:.9f}",
            "weight=WEIGHT; split modules require SP1WEIGHT or SP2WEIGHT",
            "replicate_weights=REPWEIGHT1-REPWEIGHT160",
            "variance=4/160 times sum of squared replicate deviations",
            "record_unit=relational; CONTROL joins child tables to household",
            "project_table_has_no_source-defined unique row key",
            "imputation=J-prefixed flags; not applicable is distinct from not reported",
            "limitation=2023 national unit response rate was 59.2 percent",
        ]
        return ValidationResult(
            passed=all(checks.values()),
            checks=checks,
            notes=notes,
        )

    def normalize_release(
        self,
        run_root: Path,
        release: DiscoveredRelease,
        artifacts: list[StoredArtifact],
    ) -> list[Path]:
        del release
        extracted = _extracted_tables(run_root, artifacts, self.definition)
        normalized = run_root / "normalized"
        normalized.mkdir(parents=True, exist_ok=True)
        outputs: list[Path] = []
        for table_name in self.definition.tables:
            source = extracted[table_name]
            output = normalized / f"{table_name}.parquet"
            _csv_to_parquet(source, output)
            outputs.append(output)
        return outputs


def _load_definition(config_dir: Path) -> AHSReleaseDefinition:
    return AHSReleaseDefinition.model_validate(
        yaml.safe_load((config_dir / "2023.yaml").read_text())
    )


def _extract_links(html: str, base_url: str) -> list[tuple[str, str]]:
    links: list[tuple[str, str]] = []
    for href, text in re.findall(
        r'<a[^>]+href=["\']([^"\']+)["\'][^>]*>(.*?)</a>',
        html,
        flags=re.IGNORECASE | re.DOTALL,
    ):
        label = re.sub(r"<[^>]+>", " ", text)
        label = re.sub(r"\s+", " ", label).strip()
        links.append((urljoin(base_url, href), label))
    return links


def _resolve_page_artifacts(
    links: list[tuple[str, str]], *, year: int, revision: str
) -> list[DiscoveredArtifact]:
    patterns = {
        NATIONAL_CSV: rf"AHS {year} National PUF v{re.escape(revision)} CSV\.zip$",
        VERSION_CONTROL: rf"AHS {year} National PUF Version Control\.pdf$",
        VERIFICATION_WORKBOOK: (
            rf"AHS {year} Table Specifications and PUF Estimates "
            r"for User Verification\.xlsx$"
        ),
        VALUE_LABELS: rf"AHS {year} Value Labels Package\.zip$",
    }
    found: dict[str, DiscoveredArtifact] = {}
    for url, label in links:
        filename = Path(unquote(urlparse(url).path)).name
        for role, pattern in patterns.items():
            if re.fullmatch(pattern, filename, flags=re.IGNORECASE):
                found[role] = DiscoveredArtifact(
                    role=role,
                    url=_http_url.validate_python(url),
                    filename=filename,
                    documentation=role == VERSION_CONTROL,
                    link_text=label,
                )
    missing = set(patterns) - set(found)
    if missing:
        raise RuntimeError(f"AHS page missing expected artifacts: {sorted(missing)}")
    return [found[role] for role in patterns]


def _extracted_tables(
    release_root: Path,
    artifacts: list[StoredArtifact],
    definition: AHSReleaseDefinition,
) -> dict[str, Path]:
    data = next(
        (item for item in artifacts if _role_value(item.role) == NATIONAL_CSV),
        None,
    )
    if data is None:
        return {}
    expected = {contract.filename: name for name, contract in definition.tables.items()}
    result: dict[str, Path] = {}
    for relative in data.extracted_files:
        path = release_root / relative
        table_name = expected.get(path.name)
        if table_name:
            result[table_name] = path
    return result


def _csv_header(path: Path) -> list[str]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return next(csv.reader(handle))


def _validate_relational_table(
    path: Path,
    contract: AHSTableContract,
) -> tuple[int, set[tuple[str, ...]], set[str], bool, bool]:
    columns = contract.key or ["CONTROL"]
    rows = 0
    keys: set[tuple[str, ...]] = set()
    controls: set[str] = set()
    duplicates = False
    null_keys = False
    for chunk in pd.read_csv(path, usecols=columns, chunksize=20_000, low_memory=False):
        rows += len(chunk)
        null_keys |= bool(chunk[columns].isna().any(axis=1).any())
        controls.update(chunk["CONTROL"].astype(str))
        if contract.key:
            for values in chunk[contract.key].itertuples(index=False, name=None):
                key = tuple(str(value) for value in values)
                if key in keys:
                    duplicates = True
                keys.add(key)
    return rows, keys, controls, duplicates, null_keys


def _normalize_code(value: object) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip().strip("'")
    return None if normalized.lower() in {"nan", "<na>"} else normalized


def _role_value(role: object) -> str:
    return str(getattr(role, "value", role))


def _artifact_path(
    release_root: Path,
    artifacts: list[StoredArtifact],
    role: str,
) -> Path:
    matches = [item for item in artifacts if _role_value(item.role) == role]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one AHS {role} artifact, found {len(matches)}")
    return release_root / matches[0].relative_path


def _csv_to_parquet(csv_path: Path, output: Path) -> None:
    reader = pacsv.open_csv(
        csv_path,
        read_options=pacsv.ReadOptions(block_size=32 * 1024 * 1024, use_threads=True),
        convert_options=pacsv.ConvertOptions(strings_can_be_null=True),
    )
    with pq.ParquetWriter(output, reader.schema, compression="zstd") as writer:
        for batch in reader:
            writer.write_batch(batch)
